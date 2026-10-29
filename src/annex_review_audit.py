"""Explicit one-opportunity read audit, separate from ordinary scan approval.

The regular reader keeps its three-document limit and supported formats.
This CLI inspects at most seven remaining files from the exact observed sample,
with bounded bytes, expansion, text and subprocess resources. It never updates
an opportunity, sends messages or chooses an authoritative document revision.
"""
import hashlib
import io
import json
import re
import resource
import subprocess
import tempfile
import zipfile
from xml.etree import ElementTree as ET

from src import document_parser as reader
from src.ted_annex_probe import PUBLICATION, PORTAL
from src.ted_bulk import fetch_daily_notices

MAX_REMAINING = 7
MAX_EXPANDED = 64 * 1024 * 1024
MAX_CELLS = 10000
MAX_TEXT_BYTES = 1024 * 1024


def _child_limits():
    resource.setrlimit(resource.RLIMIT_CPU, (8, 8))
    resource.setrlimit(resource.RLIMIT_AS, (256 * 1024 * 1024, 256 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_TEXT_BYTES, MAX_TEXT_BYTES))


def legacy_word(data):
    if not data.startswith(bytes.fromhex('d0cf11e0a1b11ae1')):
        raise ValueError('legacy_word_magic_missing')
    with tempfile.TemporaryFile() as output:
        result = subprocess.run(['antiword', '-m', 'UTF-8.txt', '-'], input=data,
                                stdout=output, stderr=subprocess.DEVNULL,
                                timeout=10, preexec_fn=_child_limits, check=False)
        if result.returncode != 0:
            raise ValueError('legacy_word_reader_failed')
        output.seek(0)
        raw = output.read(MAX_TEXT_BYTES + 1)
    if len(raw) > MAX_TEXT_BYTES:
        raise ValueError('legacy_word_text_limit')
    text = raw.decode('utf-8', errors='strict')
    return {'text_chars': min(len(text), reader.MAX_EXTRACTED_CHARS),
            'text_truncated': len(text) > reader.MAX_EXTRACTED_CHARS,
            'formula_cells': 0}


def spreadsheet(data):
    # Never extract archive paths, evaluate formulas or follow external links.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > 2048 or sum(m.file_size for m in members) > MAX_EXPANDED:
            raise ValueError('spreadsheet_archive_limit')
        if any(m.flag_bits & 1 or m.filename.startswith('/') or '..' in m.filename.split('/') for m in members):
            raise ValueError('spreadsheet_archive_unsafe')
        if len({m.filename for m in members}) != len(members):
            raise ValueError('spreadsheet_duplicate_member')
        if any(m.filename.lower().endswith('vbaproject.bin') for m in members):
            raise ValueError('spreadsheet_macro_member')
        for member in members:
            if member.filename.endswith(('.xml', '.rels')):
                if member.file_size > reader.MAX_DOCUMENT_BYTES:
                    raise ValueError('spreadsheet_xml_limit')
                xml = archive.read(member)
                if re.search(br'<!\s*(?:DOCTYPE|ENTITY)', xml, re.I):
                    raise ValueError('spreadsheet_xml_declaration')
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    count = chars = formulas = 0
    truncated = False
    try:
        for sheet in workbook.worksheets:
            # Worksheet dimensions are untrusted metadata; bound iteration too.
            row_limit = min(sheet.max_row or 10000, 10000)
            col_limit = min(sheet.max_column or 100, 100)
            dimensions_clipped = (sheet.max_row or 10000) > row_limit or (sheet.max_column or 100) > col_limit
            for row in sheet.iter_rows(min_row=1, max_row=row_limit, max_col=col_limit):
                for cell in row:
                    count += 1
                    if count > MAX_CELLS:
                        truncated = True
                        break
                    if cell.value is not None:
                        formulas += int(cell.data_type == 'f')
                        chars += len(str(cell.value))
                    if chars > reader.MAX_EXTRACTED_CHARS:
                        truncated = True
                        break
                if truncated:
                    break
            if dimensions_clipped:
                truncated = True
            if truncated:
                break
    finally:
        workbook.close()
    return {'text_chars': min(chars, reader.MAX_EXTRACTED_CHARS),
            'text_truncated': truncated, 'formula_cells': formulas}


def audit_bytes(kind, data, content_type, final_url):
    if not data or len(data) > reader.MAX_DOCUMENT_BYTES:
        raise ValueError('audit_document_byte_limit')
    if kind == 'doc':
        return legacy_word(data)
    if kind == 'xlsx':
        return spreadsheet(data)
    text, actual_kind = reader._parse(data, content_type, final_url)
    if kind not in {'pdf', 'docx', 'txt'}:
        raise ValueError('audit_format_not_supported')
    return {'text_chars': len(text), 'text_truncated': len(text) >= reader.MAX_EXTRACTED_CHARS,
            'formula_cells': 0}


def main():
    notice = fetch_daily_notices('197/2026', [PUBLICATION])[PUBLICATION]
    references = [(child.text or '').strip() for node in ET.fromstring(notice.data).iter()
                  if node.tag.rsplit('}', 1)[-1] == 'CallForTendersDocumentReference'
                  for child in node.iter() if child.tag.rsplit('}', 1)[-1] == 'URI']
    if PORTAL not in references:
        raise ValueError('audit_notice_portal_mismatch')
    row = reader.enrich_opportunity({'document_url': PORTAL})
    inventory = row['document_extraction'].get('attachment_inventory', [])
    pending = [entry for entry in inventory if entry.get('status') != 'parsed']
    if len(inventory) != 10 or len(pending) > MAX_REMAINING:
        raise ValueError('audit_sample_scope_changed')
    records = []
    for index, entry in enumerate(pending, 1):
        kind = entry.get('kind')
        record = {'index': index, 'format': kind, 'status': 'unread'}
        try:
            # Same public-DNS/redirect and per-document byte guards as the scan.
            data, content_type, final_url = reader._download(entry['url'])
            if reader.same_origin(entry['url'], final_url) is not True:
                raise ValueError('audit_attachment_origin_changed')
            record.update(audit_bytes(kind, data, content_type, final_url))
            if record['text_chars'] < 20:
                raise ValueError('audit_extractable_text_missing')
            record.update(status='read_for_review', sha256=hashlib.sha256(data).hexdigest(), bytes=len(data))
        except (ValueError, OSError, RuntimeError, subprocess.SubprocessError, zipfile.BadZipFile):
            record['status'] = 'manual_review_required'
        records.append(record)
    # Safe metadata only: no filenames, links, source text, counterparties or cells.
    print('BOUNDED_ANNEX_AUDIT', json.dumps({'publication': PUBLICATION,
        'ordinary_reader_parsed': len(inventory) - len(pending), 'audited_remaining': records,
        'ordinary_reader_limits_unchanged': True, 'all_mandatory_requirements_reviewed': False,
        'applicable_revision_selected': False, 'supplier_sourcing_ready': False,
        'commercial_approval': False}, sort_keys=True), flush=True)


if __name__ == '__main__':
    main()
